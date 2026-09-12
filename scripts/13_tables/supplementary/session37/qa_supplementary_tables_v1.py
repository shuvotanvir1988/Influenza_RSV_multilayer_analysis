from pathlib import Path
import hashlib
import json
import math

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path.home() / "Influenza_RSV_Project"

MANIFEST = (
    ROOT
    / "results/session37_supplementary_tables/"
      "Session37_Supplementary_Table_Manifest_v1.0.tsv"
)

WORKBOOK_DIR = (
    ROOT
    / "results/session37_supplementary_tables/workbooks"
)

OUTDIR = (
    ROOT
    / "results/session37_supplementary_tables"
)

QCDIR = OUTDIR / "qa"
QCDIR.mkdir(parents=True, exist_ok=True)


S1_DISCOVERY_COLUMNS = [
    "canonical_sample_id",
    "dataset_id",
    "gse_accession",
    "platform_id",
    "geo_accession",
    "age",
    "gender",
    "sample_group",
    "tissue",
    "age_months",
    "diagnosis",
    "harmonized_group",
    "pathogen",
    "infection_phase",
    "is_control",
    "sex_standardized",
    "tissue_standardized",
    "age_months_harmonized",
    "age_unit_status",
    "age_harmonization_note",
    "eligible_primary_infection_vs_control",
    "eligible_direct_influenza_vs_rsv",
    "eligible_rsv_acute_vs_control",
    "eligible_rsv_recovery_analysis",
    "eligible_age_adjusted_model",
    "qc_unclassified_group",
    "qc_unknown_sex",
    "qc_unknown_tissue",
    "qc_duplicate_geo_accession",
]


EXPECTED_KEY_COUNTS = {
    ("S1", "Discovery_Cohort"): (241, 29),
    ("S1", "GSE155925_Eligibility"): (64, 12),
    ("S1", "SIG_Cohort"): (208, 8),
    ("S1", "GSE283746_Cohort"): (370, 14),

    ("S2", "Frozen170_Architecture"): (170, 9),
    ("S2", "DE_Robustness"): (5, 3),

    ("S3", "Domain_Summary"): (7, 13),
    ("S3", "Antiviral_Programs"): (5, 4),
    ("S3", "Leading_Edge"): (13, 3),
    ("S3", "Hallmark"): (25, 12),
    ("S3", "Reactome"): (10, 12),

    ("S4", "Pathway_Adjustment"): (6, 17),
    ("S4", "Gene_Adjustment"): (8, 12),
    ("S4", "Signature_Adjustment"): (3, 12),

    ("S5", "Regulatory37"): (37, 70),
    ("S5", "PathogenBiased13"): (13, 71),
    ("S5", "Bulk_Validation"): (37, 79),
    ("S5", "SingleCell_Validation"): (407, 13),

    ("S6", "Gene_Level_170"): (170, 59),
    ("S6", "Season_Summary"): (16, 15),
    ("S6", "CrossSeason_Summary"): (4, 8),

    ("S7", "GSE155925_Genes"): (170, 29),
    ("S7", "GSE155925_Summary"): (4, 16),
    ("S7", "scRNA_CellTypes"): (44, 21),
    ("S7", "scRNA_ISG_State"): (8, 12),

    ("S8", "Gene_Evidence"): (170, 53),
    ("S8", "Functional_Class"): (170, 63),
    ("S8", "Direction_Aware"): (170, 63),
    ("S8", "Pathways"): (168, 19),
    ("S8", "Domains"): (21, 13),

    ("S9", "Assay_Mapping"): (185, 15),
    ("S9", "Gene_Validation"): (84, 15),
    ("S9", "Domain_Validation"): (7, 10),
    ("S9", "Regulatory_Validation"): (37, 19),

    ("S10", "Integrated170"): (170, 107),
    ("S10", "Selected_Genes"): (17, 108),
    ("S10", "Architecture_Summary"): (8, 5),
}


EXPECTED_SHEETS = {
    "S1": [
        "README",
        "Discovery_Cohort",
        "GSE155925_Eligibility",
        "SIG_Cohort",
        "GSE283746_Cohort",
    ],
    "S2": [
        "README",
        "Frozen170_Architecture",
        "DE_Robustness",
    ],
    "S3": [
        "README",
        "Domain_Summary",
        "Antiviral_Programs",
        "Leading_Edge",
        "Hallmark",
        "Reactome",
    ],
    "S4": [
        "README",
        "Pathway_Adjustment",
        "Gene_Adjustment",
        "Signature_Adjustment",
    ],
    "S5": [
        "README",
        "Regulatory37",
        "PathogenBiased13",
        "Bulk_Validation",
        "SingleCell_Validation",
    ],
    "S6": [
        "README",
        "Gene_Level_170",
        "Season_Summary",
        "CrossSeason_Summary",
    ],
    "S7": [
        "README",
        "GSE155925_Genes",
        "GSE155925_Summary",
        "scRNA_CellTypes",
        "scRNA_ISG_State",
    ],
    "S8": [
        "README",
        "Gene_Evidence",
        "Functional_Class",
        "Direction_Aware",
        "Pathways",
        "Domains",
    ],
    "S9": [
        "README",
        "Assay_Mapping",
        "Gene_Validation",
        "Domain_Validation",
        "Regulatory_Validation",
    ],
    "S10": [
        "README",
        "Integrated170",
        "Selected_Genes",
        "Architecture_Summary",
    ],
}


def sha256_file(path):
    h = hashlib.sha256()

    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)

    return h.hexdigest()


def read_source(path):
    suffix = path.suffix.lower()

    if suffix == ".tsv":
        return pd.read_csv(path, sep="\t", low_memory=False)

    if suffix == ".csv":
        return pd.read_csv(path, low_memory=False)

    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path)

    raise ValueError(f"Unsupported format: {path}")


def normalize_value(v):
    """
    Normalize source/workbook scalar values for semantic comparison.

    Excel may read integral numeric values such as 4.0 as integer 4.
    These are scientifically identical and are normalized to float
    before comparison.
    """

    if v is None:
        return None

    try:
        if pd.isna(v):
            return None
    except Exception:
        pass

    if isinstance(v, (bool, np.bool_)):
        return bool(v)

    if isinstance(v, (int, np.integer)):
        return float(v)

    if isinstance(v, (float, np.floating)):
        fv = float(v)

        if math.isnan(fv):
            return None

        return fv

    if isinstance(v, np.generic):
        try:
            v = v.item()
        except Exception:
            pass

    s = str(v)

    if s == "True":
        return True

    if s == "False":
        return False

    return s


def values_equal(a, b):
    a = normalize_value(a)
    b = normalize_value(b)

    if a is None and b is None:
        return True

    if isinstance(a, bool) or isinstance(b, bool):
        return a == b

    if (
        isinstance(a, (int, float))
        and isinstance(b, (int, float))
    ):
        return math.isclose(
            float(a),
            float(b),
            rel_tol=1e-12,
            abs_tol=1e-12,
        )

    return str(a) == str(b)


def worksheet_dataframe(ws):
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        return pd.DataFrame()

    header = [
        str(x) if x is not None else ""
        for x in rows[0]
    ]

    return pd.DataFrame(
        rows[1:],
        columns=header
    )


manifest = pd.read_csv(
    MANIFEST,
    sep="\t",
    dtype=str,
)

manifest_submission = manifest[
    manifest["inclusion"] == "INCLUDE"
].copy()


qa_rows = []
mismatch_rows = []
readme_rows = []
overall_failures = []


print("=== SESSION 37 SUPPLEMENTARY TABLE QA ===")
print(f"Submission worksheet records: {len(manifest_submission)}")
print()


for table_id in [f"S{i}" for i in range(1, 11)]:

    workbook_path = (
        WORKBOOK_DIR
        / f"Supplementary_Table_{table_id}.xlsx"
    )

    print("=" * 88)
    print(table_id, workbook_path.name)
    print("=" * 88)

    if not workbook_path.exists():
        overall_failures.append(
            f"{table_id}: workbook missing"
        )
        print("FAIL: workbook missing")
        continue

    wb = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
    )

    expected = EXPECTED_SHEETS[table_id]
    observed = wb.sheetnames

    sheet_order_ok = observed == expected

    print(
        "Sheet structure:",
        "PASS" if sheet_order_ok else "FAIL"
    )

    if not sheet_order_ok:
        overall_failures.append(
            f"{table_id}: worksheet structure mismatch"
        )

    readme = wb["README"]

    if readme.freeze_panes != "A7":
        overall_failures.append(
            f"{table_id}: README freeze pane mismatch"
        )

    if readme["A1"].value != f"Supplementary Table {table_id[1:]}":
        overall_failures.append(
            f"{table_id}: README title mismatch"
        )

    manifest_rows = manifest_submission[
        manifest_submission["supp_table"] == table_id
    ]

    readme_source_map = {}

    for r in range(7, readme.max_row + 1):
        worksheet = readme.cell(r, 1).value
        source = readme.cell(r, 2).value
        digest = readme.cell(r, 3).value

        if worksheet and source:
            readme_source_map[str(worksheet)] = {
                "source": str(source),
                "sha256": str(digest),
            }

    table_readme_ok = []

    for _, rec in manifest_rows.iterrows():

        ws_name = rec["worksheet"]
        source_rel = rec["source_path"]
        source_path = ROOT / source_rel
        source_sha = sha256_file(source_path)

        embedded = readme_source_map.get(ws_name)

        map_ok = (
            embedded is not None
            and embedded["source"] == source_rel
            and embedded["sha256"] == source_sha
        )

        table_readme_ok.append(map_ok)

        readme_rows.append({
            "table": table_id,
            "worksheet": ws_name,
            "source_path": source_rel,
            "expected_sha256": source_sha,
            "embedded_source": (
                embedded["source"] if embedded else ""
            ),
            "embedded_sha256": (
                embedded["sha256"] if embedded else ""
            ),
            "readme_mapping_ok": map_ok,
        })

        if not map_ok:
            overall_failures.append(
                f"{table_id}/{ws_name}: README provenance mismatch"
            )

    print(
        "README provenance:",
        "PASS" if all(table_readme_ok) else "FAIL"
    )

    for _, rec in manifest_rows.iterrows():

        ws_name = rec["worksheet"]
        source_path = ROOT / rec["source_path"]

        source_df = read_source(source_path)

        if (
            table_id == "S1"
            and ws_name == "Discovery_Cohort"
        ):
            source_df = source_df[
                S1_DISCOVERY_COLUMNS
            ].copy()

        ws = wb[ws_name]
        workbook_df = worksheet_dataframe(ws)

        expected_shape = EXPECTED_KEY_COUNTS[
            (table_id, ws_name)
        ]

        shape_ok = (
            workbook_df.shape == expected_shape
            and source_df.shape == expected_shape
        )

        columns_ok = (
            list(workbook_df.columns)
            == list(source_df.columns)
        )

        freeze_ok = ws.freeze_panes == "A2"

        filter_ok = (
            ws.auto_filter.ref is not None
            and str(ws.auto_filter.ref) != ""
        )

        excel_table_count = len(ws.tables)
        excel_table_ok = excel_table_count == 1

        comparison_ok = True
        mismatch_count = 0

        if (
            columns_ok
            and source_df.shape == workbook_df.shape
        ):

            for r in range(len(source_df)):
                for c in range(len(source_df.columns)):

                    a = source_df.iat[r, c]
                    b = workbook_df.iat[r, c]

                    if not values_equal(a, b):

                        comparison_ok = False
                        mismatch_count += 1

                        mismatch_rows.append({
                            "table": table_id,
                            "worksheet": ws_name,
                            "row_1based_data": r + 2,
                            "column": source_df.columns[c],
                            "source_value": repr(a),
                            "workbook_value": repr(b),
                            "normalized_source": repr(
                                normalize_value(a)
                            ),
                            "normalized_workbook": repr(
                                normalize_value(b)
                            ),
                        })

        else:
            comparison_ok = False

        qa_rows.append({
            "table": table_id,
            "worksheet": ws_name,
            "source_rows": len(source_df),
            "source_cols": len(source_df.columns),
            "workbook_rows": len(workbook_df),
            "workbook_cols": len(workbook_df.columns),
            "shape_ok": shape_ok,
            "columns_ok": columns_ok,
            "values_ok": comparison_ok,
            "mismatch_count": mismatch_count,
            "freeze_panes_ok": freeze_ok,
            "autofilter_ok": filter_ok,
            "excel_table_ok": excel_table_ok,
            "excel_table_count": excel_table_count,
        })

        local_ok = all([
            shape_ok,
            columns_ok,
            comparison_ok,
            freeze_ok,
            filter_ok,
            excel_table_ok,
        ])

        print(
            f"{ws_name:<25} "
            f"{workbook_df.shape[0]:>4} x "
            f"{workbook_df.shape[1]:>3} | "
            f"{'PASS' if local_ok else 'FAIL'}"
        )

        if not local_ok:
            overall_failures.append(
                f"{table_id}/{ws_name}: worksheet QA failed"
            )

    formula_cells = 0
    formula_error_literals = []

    error_tokens = {
        "#REF!",
        "#DIV/0!",
        "#VALUE!",
        "#NAME?",
        "#N/A",
        "#NUM!",
        "#NULL!",
    }

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:

                v = cell.value

                if isinstance(v, str) and v.startswith("="):
                    formula_cells += 1

                if isinstance(v, str) and v in error_tokens:
                    formula_error_literals.append(
                        f"{ws.title}!{cell.coordinate}:{v}"
                    )

    print(
        f"Formula cells: {formula_cells} | "
        f"{'PASS' if formula_cells == 0 else 'REVIEW'}"
    )

    print(
        f"Excel error literals: "
        f"{len(formula_error_literals)} | "
        f"{'PASS' if len(formula_error_literals) == 0 else 'FAIL'}"
    )

    if formula_cells != 0:
        overall_failures.append(
            f"{table_id}: unexpected formulas found"
        )

    if formula_error_literals:
        overall_failures.append(
            f"{table_id}: Excel error literals found"
        )

    wb.close()
    print()


qa_df = pd.DataFrame(qa_rows)
mismatch_df = pd.DataFrame(mismatch_rows)
readme_df = pd.DataFrame(readme_rows)


qa_out = QCDIR / "Session37_workbook_QA_v1.0.tsv"

mismatch_out = (
    QCDIR
    / "Session37_value_mismatches_v1.0.tsv"
)

readme_out = (
    QCDIR
    / "Session37_README_provenance_QA_v1.0.tsv"
)


qa_df.to_csv(
    qa_out,
    sep="\t",
    index=False
)

mismatch_df.to_csv(
    mismatch_out,
    sep="\t",
    index=False
)

readme_df.to_csv(
    readme_out,
    sep="\t",
    index=False
)


assertions = []


def add_assertion(name, observed, expected):
    ok = observed == expected

    assertions.append({
        "assertion": name,
        "observed": observed,
        "expected": expected,
        "pass": ok,
    })

    if not ok:
        overall_failures.append(
            f"Assertion failed: {name}"
        )


def get_sheet_df(table_id, sheet_name):

    p = (
        WORKBOOK_DIR
        / f"Supplementary_Table_{table_id}.xlsx"
    )

    wb = load_workbook(
        p,
        read_only=True,
        data_only=False
    )

    ws = wb[sheet_name]
    df = worksheet_dataframe(ws)

    wb.close()

    return df


s2 = get_sheet_df(
    "S2",
    "Frozen170_Architecture"
)

add_assertion(
    "S2 frozen architecture gene count",
    len(s2),
    170
)


s5_37 = get_sheet_df(
    "S5",
    "Regulatory37"
)

s5_13 = get_sheet_df(
    "S5",
    "PathogenBiased13"
)

add_assertion(
    "S5 high-confidence regulatory drivers",
    len(s5_37),
    37
)

add_assertion(
    "S5 pathogen-biased drivers",
    len(s5_13),
    13
)


s6 = get_sheet_df(
    "S6",
    "Gene_Level_170"
)

add_assertion(
    "S6 multiseason frozen genes",
    len(s6),
    170
)


s7 = get_sheet_df(
    "S7",
    "GSE155925_Genes"
)

add_assertion(
    "S7 GSE155925 frozen genes",
    len(s7),
    170
)


s8_gene = get_sheet_df(
    "S8",
    "Gene_Evidence"
)

s8_func = get_sheet_df(
    "S8",
    "Functional_Class"
)

s8_dir = get_sheet_df(
    "S8",
    "Direction_Aware"
)

add_assertion(
    "S8 CRISPR gene evidence genes",
    len(s8_gene),
    170
)

add_assertion(
    "S8 CRISPR functional-class genes",
    len(s8_func),
    170
)

add_assertion(
    "S8 CRISPR direction-aware genes",
    len(s8_dir),
    170
)


s9 = get_sheet_df(
    "S9",
    "Gene_Validation"
)

add_assertion(
    "S9 measurable protein genes",
    len(s9),
    84
)


s10 = get_sheet_df(
    "S10",
    "Integrated170"
)

s10sel = get_sheet_df(
    "S10",
    "Selected_Genes"
)

add_assertion(
    "S10 integrated genes",
    len(s10),
    170
)

add_assertion(
    "S10 selected multilayer genes",
    len(s10sel),
    17
)


for name, df in [
    ("S2 Frozen170", s2),
    ("S6 Gene_Level_170", s6),
    ("S7 GSE155925_Genes", s7),
    ("S8 Gene_Evidence", s8_gene),
    ("S9 Gene_Validation", s9),
    ("S10 Integrated170", s10),
]:

    if "gene_symbol" in df.columns:

        duplicates = int(
            df["gene_symbol"]
            .dropna()
            .duplicated()
            .sum()
        )

        add_assertion(
            f"{name} duplicate gene symbols",
            duplicates,
            0
        )


assertions_df = pd.DataFrame(assertions)

assertions_out = (
    QCDIR
    / "Session37_biological_cardinality_assertions_v1.0.tsv"
)

assertions_df.to_csv(
    assertions_out,
    sep="\t",
    index=False
)


checksum_rows = []

for table_id in [f"S{i}" for i in range(1, 11)]:

    p = (
        WORKBOOK_DIR
        / f"Supplementary_Table_{table_id}.xlsx"
    )

    checksum_rows.append({
        "table": table_id,
        "path": str(p.relative_to(ROOT)),
        "sha256": sha256_file(p),
        "size_bytes": p.stat().st_size,
    })


checksum_df = pd.DataFrame(checksum_rows)

checksum_out = (
    QCDIR
    / "Session37_workbook_SHA256_QA_v1.0.tsv"
)

checksum_df.to_csv(
    checksum_out,
    sep="\t",
    index=False
)


worksheet_failures = qa_df[
    ~(
        qa_df["shape_ok"]
        & qa_df["columns_ok"]
        & qa_df["values_ok"]
        & qa_df["freeze_panes_ok"]
        & qa_df["autofilter_ok"]
        & qa_df["excel_table_ok"]
    )
]

readme_failures = readme_df[
    readme_df["readme_mapping_ok"] != True
]

assertion_failures = assertions_df[
    assertions_df["pass"] != True
]


total_value_mismatches = int(
    qa_df["mismatch_count"].sum()
)


summary = {
    "submission_data_worksheets_expected": 37,
    "submission_data_worksheets_tested": int(len(qa_df)),
    "worksheet_failures": int(len(worksheet_failures)),
    "value_mismatch_records": total_value_mismatches,
    "readme_provenance_failures": int(len(readme_failures)),
    "biological_assertion_failures": int(
        len(assertion_failures)
    ),
    "overall_pass": (
        len(worksheet_failures) == 0
        and total_value_mismatches == 0
        and len(readme_failures) == 0
        and len(assertion_failures) == 0
        and len(overall_failures) == 0
    ),
}


summary_path = (
    QCDIR
    / "Session37_QA_SUMMARY_v1.0.json"
)

with open(summary_path, "w") as f:
    json.dump(
        summary,
        f,
        indent=2
    )


print("=" * 88)
print("FINAL QA SUMMARY")
print("=" * 88)

for k, v in summary.items():
    print(f"{k}: {v}")

print()

if summary["overall_pass"]:
    print("SESSION37_QA_STATUS: PASS")

else:
    print("SESSION37_QA_STATUS: FAIL")
    print()
    print("Failure details:")

    for item in sorted(set(overall_failures)):
        print(" -", item)


print()
print("QA outputs:")
print(" ", qa_out.relative_to(ROOT))
print(" ", mismatch_out.relative_to(ROOT))
print(" ", readme_out.relative_to(ROOT))
print(" ", assertions_out.relative_to(ROOT))
print(" ", checksum_out.relative_to(ROOT))
print(" ", summary_path.relative_to(ROOT))
