from pathlib import Path
import hashlib
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path.home() / "Influenza_RSV_Project"
MANIFEST = (
    ROOT
    / "results/session37_supplementary_tables/"
      "Session37_Supplementary_Table_Manifest_v1.0.tsv"
)
OUTDIR = (
    ROOT
    / "results/session37_supplementary_tables/workbooks"
)

OUTDIR.mkdir(parents=True, exist_ok=True)


TABLE_TITLES = {
    "S1": "Study cohorts, sample characteristics, and analysis eligibility",
    "S2": "Influenza differential-expression and replicated gene architecture",
    "S3": "Influenza pathway and leading-edge analyses",
    "S4": "Immune-cell composition and composition-adjustment analyses",
    "S5": "Regulatory architecture and validation of candidate drivers",
    "S6": "Four-season independent validation of the influenza transcriptional architecture",
    "S7": "Independent bulk and single-cell transcriptomic validation",
    "S8": "CRISPR functional validation and pathway convergence",
    "S9": "Proteomic validation of the influenza transcriptional architecture",
    "S10": "Integrated multilayer evidence for the frozen 170-gene architecture",
}


# ------------------------------------------------------------
# Discovery metadata:
# deliberately remove GEO administrative/contact/protocol fields.
# ------------------------------------------------------------

S1_DISCOVERY_COLUMNS = [
    "canonical_sample_id",
    "dataset_id",
    "gse_accession",
    "platform_id",
    "geo_accession",
    "age",
    "gender",
    "sample_group",
    "tissue",
    "age_months",
    "diagnosis",
    "harmonized_group",
    "pathogen",
    "infection_phase",
    "is_control",
    "sex_standardized",
    "tissue_standardized",
    "age_months_harmonized",
    "age_unit_status",
    "age_harmonization_note",
    "eligible_primary_infection_vs_control",
    "eligible_direct_influenza_vs_rsv",
    "eligible_rsv_acute_vs_control",
    "eligible_rsv_recovery_analysis",
    "eligible_age_adjusted_model",
    "qc_unclassified_group",
    "qc_unknown_sex",
    "qc_unknown_tissue",
    "qc_duplicate_geo_accession",
]


README_DESCRIPTIONS = {
    "S1": (
        "Sample-level cohort, phenotype, eligibility, and validation-cohort "
        "information used in the study. Discovery GEO administrative/contact "
        "metadata are intentionally excluded from the submission worksheet."
    ),
    "S2": (
        "Frozen 170-gene replicated architecture and differential-expression "
        "robustness analyses."
    ),
    "S3": (
        "Preregistered pathway-domain summaries, direct antiviral programs, "
        "leading-edge recurrence, and representative Hallmark/Reactome results."
    ),
    "S4": (
        "Sensitivity analyses evaluating whether influenza-associated pathway, "
        "gene, and signature effects persist after adjustment for estimated "
        "immune-cell composition."
    ),
    "S5": (
        "Frozen regulatory-driver architecture together with independent bulk "
        "RNA-seq and single-cell validation."
    ),
    "S6": (
        "Gene-level and summary validation of the frozen transcriptional "
        "architecture across four independent influenza seasons."
    ),
    "S7": (
        "Independent bulk RNA-seq and single-cell transcriptomic validation "
        "of the frozen gene architecture."
    ),
    "S8": (
        "Gene-level, functional-classification, direction-aware, pathway, and "
        "domain-level CRISPR evidence."
    ),
    "S9": (
        "Influenza SomaScan assay mapping and gene-, domain-, and regulatory-"
        "program proteomic validation."
    ),
    "S10": (
        "Integrated transcriptomic, proteomic, CRISPR, and regulatory evidence "
        "for the frozen 170-gene architecture."
    ),
}


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_source(path):
    suffix = path.suffix.lower()

    if suffix == ".tsv":
        return pd.read_csv(path, sep="\t", low_memory=False)

    if suffix == ".csv":
        return pd.read_csv(path, low_memory=False)

    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path)

    raise ValueError(f"Unsupported source format: {path}")


def clean_value(v):
    if pd.isna(v):
        return None

    # Convert numpy scalar values into native Python values.
    if hasattr(v, "item"):
        try:
            return v.item()
        except Exception:
            pass

    return v


def write_dataframe(ws, df):
    # Header
    for col_idx, col in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=str(col))

    # Data
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=clean_value(value)
            )


def style_data_sheet(ws):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    thin_gray = Side(
        style="thin",
        color="D9D9D9"
    )

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = Border(bottom=thin_gray)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    # Sensible column widths with a hard upper bound.
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)

        values = []
        for row_idx in range(1, min(ws.max_row, 250) + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                values.append(len(str(v)))

        max_len = max(values) if values else 8

        if col_idx == 1:
            width = min(max(max_len + 2, 14), 32)
        else:
            width = min(max(max_len + 2, 10), 28)

        ws.column_dimensions[letter].width = width

    ws.row_dimensions[1].height = 36

    # Number formats.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                name = str(
                    ws.cell(row=1, column=cell.column).value
                ).lower()

                if any(
                    key in name
                    for key in [
                        "pvalue", "_p", "p_", "fdr", "padj",
                        "permutation_p", "binomial_p"
                    ]
                ):
                    cell.number_format = "0.000E+00"

                elif any(
                    key in name
                    for key in [
                        "logfc", "log2fc", "nes", "effect",
                        "delta", "beta", "correlation",
                        "pearson", "spearman", "fraction",
                        "percent", "score"
                    ]
                ):
                    cell.number_format = "0.000"

                else:
                    cell.number_format = "0.0000"


def add_excel_table(ws, table_index):
    if ws.max_row < 2 or ws.max_column < 1:
        return

    ref = (
        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    name = f"SuppData{table_index:03d}"

    tab = Table(
        displayName=name,
        ref=ref
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    tab.tableStyleInfo = style
    ws.add_table(tab)


def make_readme(wb, table_id, manifest_rows):
    ws = wb.create_sheet("README", 0)
    ws.sheet_view.showGridLines = False

    ws["A1"] = f"Supplementary Table {table_id[1:]}"
    ws["A1"].font = Font(size=16, bold=True)

    ws["A2"] = TABLE_TITLES[table_id]
    ws["A2"].font = Font(size=13, bold=True)

    ws["A4"] = "Description"
    ws["A4"].font = Font(bold=True)

    ws["B4"] = README_DESCRIPTIONS[table_id]
    ws["B4"].alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

    ws["A6"] = "Worksheet"
    ws["B6"] = "Source file"
    ws["C6"] = "Source SHA256"
    ws["D6"] = "Notes"

    for c in ws[6]:
        c.font = Font(bold=True)
        c.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

    r = 7

    for _, rec in manifest_rows.iterrows():
        source = ROOT / rec["source_path"]

        ws.cell(r, 1, rec["worksheet"])
        ws.cell(r, 2, rec["source_path"])
        ws.cell(
            r,
            3,
            sha256_file(source) if source.exists() else "MISSING"
        )
        ws.cell(r, 4, rec["notes"])
        r += 1

    ws["A{}".format(r + 1)] = "Generation note"
    ws["A{}".format(r + 1)].font = Font(bold=True)

    ws["B{}".format(r + 1)] = (
        "Workbook generated programmatically from the frozen Session 37 "
        "manifest. Scientific values were imported from the listed source "
        "tables without recalculation."
    )
    ws["B{}".format(r + 1)].alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 68
    ws.column_dimensions["D"].width = 60

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    ws.freeze_panes = "A7"


def main():
    manifest = pd.read_csv(
        MANIFEST,
        sep="\t",
        dtype=str
    )

    # Verify all listed source files before writing anything.
    missing = []

    for _, rec in manifest.iterrows():
        source = ROOT / rec["source_path"]
        if not source.exists():
            missing.append(str(source))

    if missing:
        print("ERROR: Missing manifest source files:")
        for p in missing:
            print("  ", p)
        raise SystemExit(1)

    print("=== SESSION 37 SUPPLEMENTARY TABLE GENERATION ===")
    print(f"Manifest records: {len(manifest)}")

    generated = []

    for table_id in [f"S{i}" for i in range(1, 11)]:

        rows = manifest[
            (manifest["supp_table"] == table_id)
            & (manifest["inclusion"] == "INCLUDE")
        ].copy()

        if rows.empty:
            raise RuntimeError(
                f"No INCLUDE records for {table_id}"
            )

        wb = Workbook()

        # Remove default sheet.
        default = wb.active
        wb.remove(default)

        make_readme(
            wb,
            table_id,
            rows
        )

        table_index = 1

        for _, rec in rows.iterrows():

            source = ROOT / rec["source_path"]
            df = read_source(source)

            # S1 discovery metadata is intentionally curated.
            if (
                table_id == "S1"
                and rec["worksheet"] == "Discovery_Cohort"
            ):
                missing_cols = [
                    c
                    for c in S1_DISCOVERY_COLUMNS
                    if c not in df.columns
                ]

                if missing_cols:
                    raise RuntimeError(
                        "Missing expected S1 discovery columns: "
                        + ", ".join(missing_cols)
                    )

                df = df[S1_DISCOVERY_COLUMNS].copy()

            sheet_name = rec["worksheet"][:31]

            ws = wb.create_sheet(sheet_name)

            write_dataframe(ws, df)
            style_data_sheet(ws)
            add_excel_table(ws, table_index)

            table_index += 1

            print(
                f"{table_id:>3} | "
                f"{sheet_name:<24} | "
                f"{len(df):>5} rows x "
                f"{len(df.columns):>3} cols"
            )

        outfile = OUTDIR / f"Supplementary_Table_{table_id}.xlsx"
        wb.save(outfile)

        # Re-open to ensure workbook integrity.
        check = load_workbook(
            outfile,
            read_only=True,
            data_only=False
        )

        expected_sheets = (
            ["README"]
            + [x[:31] for x in rows["worksheet"].tolist()]
        )

        if check.sheetnames != expected_sheets:
            raise RuntimeError(
                f"{table_id}: worksheet mismatch\n"
                f"Expected: {expected_sheets}\n"
                f"Observed: {check.sheetnames}"
            )

        check.close()

        digest = sha256_file(outfile)

        generated.append(
            {
                "table": table_id,
                "file": str(outfile.relative_to(ROOT)),
                "sha256": digest,
                "worksheets": len(expected_sheets),
            }
        )

        print(
            f"    WRITTEN: {outfile.relative_to(ROOT)}"
        )

    generated_df = pd.DataFrame(generated)

    summary = (
        ROOT
        / "results/session37_supplementary_tables/"
          "Session37_generated_workbooks_v1.0.tsv"
    )

    generated_df.to_csv(
        summary,
        sep="\t",
        index=False
    )

    print()
    print("=== GENERATION COMPLETE ===")
    print(generated_df.to_string(index=False))
    print()
    print(
        "Summary:",
        summary.relative_to(ROOT)
    )


if __name__ == "__main__":
    main()
