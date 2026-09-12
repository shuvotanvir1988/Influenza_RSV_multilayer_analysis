#!/usr/bin/env python3
from pathlib import Path
import csv, hashlib
from openpyxl import load_workbook

ROOT = Path.home() / "Influenza_RSV_Project"
AUDIT = ROOT / "results/supplementary_table_inventory"
OUT = AUDIT / "manuscript_register"
OUT.mkdir(parents=True, exist_ok=True)

INV = AUDIT / "Supplementary_Tables_S1_S13_Canonical_Inventory_v1.0.tsv"
SRC = AUDIT / "Supplementary_Tables_S1_S10_Source_Manifest_Audit_v1.0.tsv"

if not INV.exists():
    raise FileNotFoundError(INV)
if not SRC.exists():
    raise FileNotFoundError(SRC)

def sha256(p):
    h=hashlib.sha256()
    with open(p,"rb") as f:
        for chunk in iter(lambda:f.read(1024*1024), b""):
            h.update(chunk)
    return h.hexdigest()

with open(INV, newline="", encoding="utf-8") as f:
    inv = {r["supp_table"]: r for r in csv.DictReader(f, delimiter="\t")}

with open(SRC, newline="", encoding="utf-8") as f:
    srcrows = list(csv.DictReader(f, delimiter="\t"))

# Conservative manuscript-facing titles based only on the frozen table organization.
titles = {
"S1":"Cohort and sample metadata",
"S2":"Frozen gene architecture and differential-expression robustness",
"S3":"Pathway and biological-program analyses",
"S4":"Cellular-composition analyses",
"S5":"Regulatory-driver analyses",
"S6":"Cross-season influenza validation",
"S7":"Independent bulk and single-cell validation",
"S8":"CRISPR functional validation",
"S9":"Proteomics validation",
"S10":"Integrated multilayer evidence",
"S11":"Evidence-weighted gene prioritization and robustness",
"S12":"Molecular–functional convergence",
"S13":"Single-cell localization and robustness",
}

register=[]
sheet_rows=[]

for n in range(1,14):
    sn=f"S{n}"
    r=inv[sn]
    wbrel=r["canonical_workbook"]
    wbpath=ROOT/wbrel
    wb=load_workbook(wbpath, read_only=True, data_only=False)

    included_sources=[]
    archived_sources=[]
    if n <= 10:
        for x in srcrows:
            if x["supp_table"] != sn:
                continue
            desc=f'{x["worksheet"]}: {x["source_path"]}'
            if x["inclusion"]=="INCLUDE":
                included_sources.append(desc)
            else:
                archived_sources.append(desc)

    for ws in wb.worksheets:
        maxr=ws.max_row or 0
        maxc=ws.max_column or 0
        sheet_rows.append({
            "supp_table":sn,
            "manuscript_title":titles[sn],
            "worksheet":ws.title,
            "rows_excluding_header":max(maxr-1,0),
            "columns":maxc,
            "canonical_workbook":wbrel,
        })

    register.append({
        "supp_table":sn,
        "manuscript_title":titles[sn],
        "canonical_workbook":wbrel,
        "workbook_sha256":sha256(wbpath),
        "worksheet_count":len(wb.sheetnames),
        "worksheets":"; ".join(wb.sheetnames),
        "session37_included_sources":" || ".join(included_sources),
        "session37_archived_only_sources":" || ".join(archived_sources),
        "status":"CANONICAL_FROZEN",
    })

reg=OUT/"Supplementary_Tables_S1_S13_Manuscript_Register_v1.0.tsv"
with open(reg,"w",newline="",encoding="utf-8") as f:
    fields=list(register[0].keys())
    w=csv.DictWriter(f,delimiter="\t",fieldnames=fields)
    w.writeheader(); w.writerows(register)

sheets=OUT/"Supplementary_Tables_S1_S13_Manuscript_Worksheet_Register_v1.0.tsv"
with open(sheets,"w",newline="",encoding="utf-8") as f:
    fields=list(sheet_rows[0].keys())
    w=csv.DictWriter(f,delimiter="\t",fieldnames=fields)
    w.writeheader(); w.writerows(sheet_rows)

# Plain-text report for easy inspection/upload.
report=OUT/"Supplementary_Tables_S1_S13_Manuscript_Register_v1.0.txt"
with open(report,"w",encoding="utf-8") as f:
    f.write("SUPPLEMENTARY TABLES S1-S13 — MANUSCRIPT REGISTER v1.0\n\n")
    for r in register:
        f.write(f'{r["supp_table"]}. {r["manuscript_title"]}\n')
        f.write(f'  Workbook: {r["canonical_workbook"]}\n')
        f.write(f'  SHA256: {r["workbook_sha256"]}\n')
        f.write(f'  Worksheets ({r["worksheet_count"]}): {r["worksheets"]}\n')
        if r["session37_archived_only_sources"]:
            f.write(f'  Archived-only Session37 source(s): {r["session37_archived_only_sources"]}\n')
        f.write("\n")

control=OUT/"Supplementary_Tables_S1_S13_Manuscript_Register_SHA256_v1.0.txt"
with open(control,"w",encoding="utf-8") as f:
    for p in [reg,sheets,report]:
        f.write(f"{sha256(p)}  {p.relative_to(ROOT)}\n")

print("=== MANUSCRIPT-FACING SUPPLEMENTARY TABLE REGISTER COMPLETE ===")
print(report)
print(reg)
print(sheets)
print(control)
print()
for r in register:
    print(f'{r["supp_table"]}: {r["manuscript_title"]} | {r["worksheet_count"]} sheets')
